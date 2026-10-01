"""Certificate run, tier C: the 16 microbiome seawater boundary fits
(m = 1, the CERTIFIED-CORNER register: a refusal there would be the
WRONG answer).

Candidate per sample: (lntheta = ln theta_ewens from the study's receipt,
d = 0 aux, q = 0 EXACTLY on its declared lower bound). bounds passed to
engine.certify so the corner path engages; verdict must come back
CERTIFIED-CORNER with the KKT multiplier sign (dNLL/dq > 0 at the lower
bound) strictly certified over the box.
"""
import json
import math
import os
import sys
import time

sys.dont_write_bytecode = True

import numpy as np

from . import DATA_DIR, SCIENCE_DIR, RECEIPTS_DIR, ident
from .oracle_corner import EtienneCornerOracle

from clinch import engine, cert

STAMP = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())

XLSX = os.path.join(
    DATA_DIR, "microbiome/sieber2019_plosbiology/"
    "pbio.3000298.s015_S5_Data_seawater_sediment.xlsx")


def load_seawater():
    """The study's own loader logic (s4_seawater_fit.load_seawater),
    replicated verbatim-in-effect; the (J, S) identity of every sample is
    asserted against the study's receipt before any certificate."""
    from openpyxl import load_workbook
    ident.log_read(XLSX, "microbiome seawater matrix (fit data of record)")
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    sheet = None
    for nm in wb.sheetnames:
        if "water" in nm.lower() or nm.lower() == "sw":
            sheet = wb[nm]
            break
    if sheet is None:
        sheet = wb[wb.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    hdr = rows[0]
    data_cols = [i for i, h in enumerate(hdr)
                 if h and all(isinstance(r[i], (int, float)) or r[i] is None
                              for r in rows[1:50])]

    def is_sample(i):
        h = str(hdr[i]).lower()
        return not any(k in h for k in ("taxon", "otu", "id", "kingdom",
                                        "phylum", "class", "order",
                                        "family", "genus", "species"))
    data_cols = [i for i in data_cols if is_sample(i)]
    M = np.array([[int(r[i] or 0) for i in data_cols] for r in rows[1:]],
                 dtype=np.int64)
    keep = M.sum(axis=0) >= 1000
    M = M[:, keep]
    M = M[M.sum(axis=1) > 0]
    names = [str(hdr[i]) for i, k in zip(data_cols, keep) if k]
    return M, names


def main():
    r_path = os.path.join(SCIENCE_DIR, "S4_SEAWATER_FITS.json")
    ident.log_read(r_path, "microbiome fit receipts (candidates of record)")
    rec = json.load(open(r_path))
    M, names = load_seawater()
    assert names == rec["samples"], "sample name mismatch vs receipt"
    assert M.sum(axis=0).tolist() == rec["J"], "J mismatch vs receipt"
    assert int(M.shape[0]) == rec["S_total"], "S_total mismatch vs receipt"
    manifest = []
    for ci, row in enumerate(rec["per_sample"]):
        tag = "corner_" + row["sample"].replace(".", "_")
        D = sorted(int(x) for x in M[:, ci] if x > 0)
        assert sum(D) == row["J"] and len(D) == row["S"], (tag, "data")
        o = EtienneCornerOracle(D, tag=tag)
        u = [math.log(row["theta_ewens"]), 0.0, 0.0]
        lo = np.array([-np.inf, -np.inf, 0.0])
        hi = np.array([np.inf, np.inf, np.inf])
        out_dir = os.path.join(RECEIPTS_DIR, tag)
        os.makedirs(out_dir, exist_ok=True)
        logf = open(os.path.join(out_dir, "LOG.txt"), "a")

        def log(s):
            logf.write(s + "\n")
            logf.flush()
        log(f"== {tag} @ {STAMP}")
        # gates: Ewens value (float register) + stationarity residual
        from flint import arb, ctx
        old = ctx.prec
        try:
            ctx.prec = 256
            x0 = ([[arb(0)], [arb(0)]], [arb(u[0])])
            nll = o.nll_ewens(x0)
            Fz, Fg = o.F(x0)
            g_th = float(Fg[0].mid())
            g_q = Fz[1][0]
        finally:
            ctx.prec = old
        vg = dict(register="receipt_float",
                  oracle_nll_ewens=float(nll.mid()),
                  receipt_lnL_ewens=row["lnL_ewens"],
                  abs_diff=abs(float(nll.mid()) + row["lnL_ewens"]),
                  bar=5e-9 * max(1.0, abs(row["lnL_ewens"])),
                  stationarity_g_lntheta_at_receipt=g_th,
                  receipt_theta_rad=row["theta_rad"])
        vg["ok"] = bool(vg["abs_diff"] <= vg["bar"])
        gates = dict(value_receipt_register=vg,
                     kkt_gradient_closed_form=dict(
                         g_q_at_q0=str(g_q),
                         K_Jm1=float(o.K_Jm1.mid()),
                         receipt_g_lnI_at_m999=row["g_lnI_at_m999"],
                         note="closed-form K_J=1, K_{J-1}=sum n_i/2 "
                              "(n_i>=2); formula battery-gated vs K_ball/"
                              "K_exact on small data"))
        t0 = time.time()
        res = engine.certify(o, u, bounds=(lo, hi), log=log)
        prov = dict(receipt=r_path, receipt_sha=ident.sha256(r_path),
                    sample=row["sample"], theta_ewens=row["theta_ewens"],
                    m="1 (boundary)  == q=0 lower bound",
                    published_register="the study's per-sample boundary fit")
        extra = dict(adapter="clinch.adapters.etienne",
                     coordinates="(lntheta | aux d | q=1/I)",
                     objective="NLL = -lnP; q=0 is the declared lower "
                               "bound (m=1, Ewens boundary)",
                     aux_note="block aux:identity is a mathematically "
                              "inert identity coordinate carried so the "
                              "masked system keeps a latent block; its "
                              "Hessian row is exactly [1] and decoupled",
                     candidate_u=u, candidate_provenance=prov,
                     run_stamp_utc=STAMP)
        cert.write_cert(res, out_dir, tag, gates=gates,
                        extra=dict(extra, product="direct_corner",
                                   kkt_leg=res.get("kkt"),
                                   active_set=res.get("active_set")))
        log(f"total {time.time()-t0:.1f}s; verdict={res['verdict']}")
        logf.close()
        manifest.append(dict(tag=tag, verdict=res["verdict"],
                             kkt=(res.get("kkt") or {}).get("ok")))
        print(tag, res["verdict"], flush=True)
    json.dump(manifest, open(os.path.join(
        RECEIPTS_DIR, "MANIFEST_TIER_C.json"), "w"), indent=1)
    ident.flush_read_log("tier_c")
    print("TIER C DONE", flush=True)


if __name__ == "__main__":
    main()
